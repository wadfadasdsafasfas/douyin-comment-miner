"""听潮 · 本地线索池（SQLite）

设计要点：
- 数据落在用户本机 ~/.tingchao/leads.db，与云端账号库分离，离线可用。
- 线索以 (平台, 视频ID, 昵称, 评论内容) 的哈希做唯一键，跨任务自动去重。
- 所有查询返回 dict / list[dict]，供 sidecar API 直接 JSON 化。
"""
from __future__ import annotations

import csv
import hashlib
import json
import sqlite3
from contextlib import contextmanager
from datetime import datetime, timedelta
from pathlib import Path

DATA_DIR = Path.home() / ".tingchao"
DB_PATH = DATA_DIR / "leads.db"

SCHEMA = """
CREATE TABLE IF NOT EXISTS leads (
    id          INTEGER PRIMARY KEY AUTOINCREMENT,
    dedup       TEXT    NOT NULL UNIQUE,
    platform    TEXT    NOT NULL DEFAULT 'douyin',
    video_id    TEXT    DEFAULT '',
    video_url   TEXT    DEFAULT '',
    nickname    TEXT    DEFAULT '',
    comment     TEXT    DEFAULT '',
    comment_time TEXT   DEFAULT '',
    region      TEXT    DEFAULT '',
    profile_url TEXT    DEFAULT '',
    dm_entry    TEXT    DEFAULT '',
    keywords    TEXT    DEFAULT '',      -- 逗号分隔的命中词
    status      TEXT    DEFAULT 'new',   -- new / contacted / replied / won / invalid
    owner       TEXT    DEFAULT '',
    note        TEXT    DEFAULT '',
    first_seen  TEXT    NOT NULL,
    last_seen   TEXT    NOT NULL,
    seen_count  INTEGER NOT NULL DEFAULT 1
);
CREATE INDEX IF NOT EXISTS idx_leads_status   ON leads(status);
CREATE INDEX IF NOT EXISTS idx_leads_platform ON leads(platform);
CREATE INDEX IF NOT EXISTS idx_leads_seen     ON leads(last_seen DESC);

CREATE TABLE IF NOT EXISTS tasks (
    id         INTEGER PRIMARY KEY AUTOINCREMENT,
    started_at TEXT NOT NULL,
    ended_at   TEXT,
    links      TEXT DEFAULT '',
    keywords   TEXT DEFAULT '',
    status     TEXT DEFAULT 'running',   -- running / done / stopped / error
    collected  INTEGER DEFAULT 0,
    hits       INTEGER DEFAULT 0,
    error      TEXT DEFAULT ''
);

CREATE TABLE IF NOT EXISTS kv (
    key   TEXT PRIMARY KEY,
    value TEXT
);
"""


@contextmanager
def db():
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    c = sqlite3.connect(DB_PATH)
    c.row_factory = sqlite3.Row
    try:
        yield c
        c.commit()
    finally:
        c.close()


def init_db() -> None:
    with db() as c:
        c.executescript(SCHEMA)


def _dedup(platform: str, video_id: str, nickname: str, comment: str) -> str:
    raw = f"{platform}|{video_id}|{nickname}|{comment}".encode("utf-8", "ignore")
    return hashlib.sha1(raw).hexdigest()


# ---------------------------------------------------------------- 写入

def upsert_lead(platform: str, row: dict, keywords_hit: list[str],
                video_url: str = "") -> bool:
    """写入一条线索；返回 True 表示新线索，False 表示已存在（仅累加 seen_count）。

    row 沿用引擎的中文字段：视频ID / 昵称 / 评论内容 / 评论时间 / 主页链接 / 私信入口
    """
    now = datetime.now().isoformat(timespec="seconds")
    nickname = (row.get("昵称") or "").strip()
    comment = (row.get("评论内容") or "").strip()
    vid = (row.get("视频ID") or "").strip()
    key = _dedup(platform, vid, nickname, comment)
    kw = ",".join(dict.fromkeys(keywords_hit))
    # 评论时间形如 "2周前 · 山东"，把地区拆出来单列存储
    raw_time = (row.get("评论时间") or "").strip()
    region = ""
    if "·" in raw_time:
        head, tail = raw_time.rsplit("·", 1)
        region, raw_time = tail.strip(), head.strip()
    with db() as c:
        is_new = c.execute("SELECT 1 FROM leads WHERE dedup=?", (key,)).fetchone() is None
        c.execute(
            """INSERT INTO leads
               (dedup,platform,video_id,video_url,nickname,comment,comment_time,region,
                profile_url,dm_entry,keywords,first_seen,last_seen)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)
               ON CONFLICT(dedup) DO UPDATE SET
                 last_seen=excluded.last_seen,
                 seen_count=seen_count+1""",
            (key, platform, vid, video_url, nickname, comment, raw_time, region,
             row.get("主页链接", ""), row.get("私信入口(主页点私信)", ""), kw, now, now))
        return is_new


def set_status(ids: list[int], status: str, owner: str = "", note: str = "") -> int:
    if not ids:
        return 0
    ph = ",".join("?" * len(ids))
    with db() as c:
        cur = c.execute(
            f"UPDATE leads SET status=?{', owner=?' if owner else ''}{', note=?' if note else ''} "
            f"WHERE id IN ({ph})",
            [status] + ([owner] if owner else []) + ([note] if note else []) + list(ids))
        return cur.rowcount


# ---------------------------------------------------------------- 查询

def list_leads(platform: str = "all", status: str = "all", q: str = "",
               page: int = 1, size: int = 20) -> dict:
    where, args = ["1=1"], []
    if platform and platform != "all":
        where.append("platform=?"); args.append(platform)
    if status and status != "all":
        where.append("status=?"); args.append(status)
    if q:
        where.append("(nickname LIKE ? OR comment LIKE ? OR keywords LIKE ?)")
        args += [f"%{q}%"] * 3
    w = " AND ".join(where)
    with db() as c:
        total = c.execute(f"SELECT COUNT(*) n FROM leads WHERE {w}", args).fetchone()["n"]
        rows = c.execute(
            f"SELECT * FROM leads WHERE {w} ORDER BY last_seen DESC LIMIT ? OFFSET ?",
            args + [size, (max(1, page) - 1) * size]).fetchall()
    return {"total": total, "page": page, "size": size,
            "items": [dict(r) for r in rows]}


def stats() -> dict:
    today = datetime.now().strftime("%Y-%m-%d")
    with db() as c:
        one = lambda sql, *a: (c.execute(sql, a).fetchone()[0] or 0)
        return {
            "total":     one("SELECT COUNT(*) FROM leads"),
            "new":       one("SELECT COUNT(*) FROM leads WHERE status='new'"),
            "contacted": one("SELECT COUNT(*) FROM leads WHERE status='contacted'"),
            "replied":   one("SELECT COUNT(*) FROM leads WHERE status='replied'"),
            "won":       one("SELECT COUNT(*) FROM leads WHERE status='won'"),
            "today":     one("SELECT COUNT(*) FROM leads WHERE first_seen LIKE ?", today + "%"),
            "users":     one("SELECT COUNT(DISTINCT platform||'|'||nickname) FROM leads"),
            "videos":    one("SELECT COUNT(DISTINCT video_id) FROM leads WHERE video_id<>''"),
            "by_platform": {r["platform"]: r["n"] for r in c.execute(
                "SELECT platform, COUNT(*) n FROM leads GROUP BY platform")},
            "by_keyword": [{"name": kw, "value": n} for kw, n in _keyword_counts(c)],
            "trend": _trend(c),
        }


def _keyword_counts(c) -> list[tuple[str, int]]:
    buckets: dict[str, int] = {}
    for r in c.execute("SELECT keywords FROM leads WHERE keywords<>''"):
        for kw in (r["keywords"] or "").split(","):
            if kw.strip():
                buckets[kw.strip()] = buckets.get(kw.strip(), 0) + 1
    return sorted(buckets.items(), key=lambda x: -x[1])[:5]


def _trend(c, days: int = 7) -> list[dict]:
    """近 N 天每日新增线索，供折线图使用。"""
    out = []
    for i in range(days - 1, -1, -1):
        d = (datetime.now() - timedelta(days=i)).strftime("%Y-%m-%d")
        out.append({"date": d[5:],
                    "leads": c.execute(
                        "SELECT COUNT(*) FROM leads WHERE first_seen LIKE ?",
                        (d + "%",)).fetchone()[0] or 0})
    return out


# ---------------------------------------------------------------- 任务记录

def task_start(links: list[str], keywords: list[str]) -> int:
    with db() as c:
        cur = c.execute("INSERT INTO tasks (started_at,links,keywords) VALUES (?,?,?)",
                        (datetime.now().isoformat(timespec="seconds"),
                         json.dumps(links, ensure_ascii=False),
                         json.dumps(keywords, ensure_ascii=False)))
        return cur.lastrowid


def task_finish(task_id: int, status: str, collected: int, hits: int, error: str = "") -> None:
    with db() as c:
        c.execute("""UPDATE tasks SET ended_at=?, status=?, collected=?, hits=?, error=? WHERE id=?""",
                  (datetime.now().isoformat(timespec="seconds"), status, collected, hits, error, task_id))


def recent_tasks(limit: int = 10) -> list[dict]:
    with db() as c:
        return [dict(r) for r in c.execute(
            "SELECT * FROM tasks ORDER BY id DESC LIMIT ?", (limit,))]


# ---------------------------------------------------------------- 导出

HEADERS = ["平台", "视频ID", "昵称", "评论内容", "命中词", "评论时间", "地区",
           "主页链接", "状态", "负责人", "备注", "首次发现", "最近发现", "出现次数"]
STATUS_CN = {"new": "待联系", "contacted": "已联系", "replied": "已回复",
             "won": "已成交", "invalid": "无效"}


def export_csv(path: str, platform: str = "all", status: str = "all", q: str = "") -> str:
    data = list_leads(platform, status, q, 1, 100000)["items"]
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(HEADERS)
        for r in data:
            w.writerow([_platform_cn(r["platform"]), r["video_id"], r["nickname"], r["comment"],
                        r["keywords"], r["comment_time"], r["region"], r["profile_url"],
                        STATUS_CN.get(r["status"], r["status"]), r["owner"], r["note"],
                        r["first_seen"], r["last_seen"], r["seen_count"]])
    return path


def export_xlsx(path: str, platform: str = "all", status: str = "all", q: str = "") -> str | None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        return None
    data = list_leads(platform, status, q, 1, 100000)["items"]
    wb = Workbook()
    ws = wb.active
    ws.title = "线索名单"
    ws.append(HEADERS)
    fill = PatternFill("solid", fgColor="FF4D4D")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r in data:
        ws.append([_platform_cn(r["platform"]), r["video_id"], r["nickname"], r["comment"],
                   r["keywords"], r["comment_time"], r["region"], r["profile_url"],
                   STATUS_CN.get(r["status"], r["status"]), r["owner"], r["note"],
                   r["first_seen"], r["last_seen"], r["seen_count"]])
    widths = [8, 20, 18, 48, 14, 12, 8, 30, 10, 10, 18, 20, 20, 10]
    for i, wdt in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i) if i <= 26 else "A" + chr(64 + i - 26)].width = wdt
    ws.freeze_panes = "A2"
    wb.save(path)
    return path


def _platform_cn(p: str) -> str:
    return {"douyin": "抖音", "xhs": "小红书", "wechat": "视频号"}.get(p, p)
